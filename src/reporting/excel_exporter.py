"""Excel export functionality."""

from pathlib import Path
from typing import List, Dict, Any
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


class ExcelExporter:
    """Export loan data and schedules to Excel."""

    @staticmethod
    def export_amortization_schedule(
        schedule: List[Dict[str, Any]],
        loan_details: Dict[str, Any],
        output_path: str
    ) -> str:
        """Export amortization schedule to Excel."""
        df = pd.DataFrame(schedule)

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            # Write schedule
            df.to_excel(writer, sheet_name='Amortization', index=False)

            # Write loan details
            loan_sheet = writer.book.create_sheet('Loan Details')
            row = 1
            for key, value in loan_details.items():
                loan_sheet[f'A{row}'] = key
                loan_sheet[f'B{row}'] = value
                row += 1

            # Style worksheets
            ExcelExporter._style_worksheet(writer.book['Amortization'])
            ExcelExporter._style_worksheet(loan_sheet)

        return output_path

    @staticmethod
    def export_comparison(
        scenarios: List[Dict[str, Any]],
        output_path: str
    ) -> str:
        """Export scenario comparison to Excel."""
        df = pd.DataFrame(scenarios)

        with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
            df.to_excel(writer, sheet_name='Comparison', index=False)
            ExcelExporter._style_worksheet(writer.book['Comparison'])

        return output_path

    @staticmethod
    def _style_worksheet(worksheet):
        """Apply styling to worksheet."""
        header_fill = PatternFill(start_color="366092", end_color="366092", fill_type="solid")
        header_font = Font(bold=True, color="FFFFFF")
        border = Border(
            left=Side(style='thin'),
            right=Side(style='thin'),
            top=Side(style='thin'),
            bottom=Side(style='thin')
        )

        # Style header row
        for cell in worksheet[1]:
            if cell.value:
                cell.fill = header_fill
                cell.font = header_font
                cell.border = border
                cell.alignment = Alignment(horizontal='center', vertical='center')

        # Style data rows
        for row in worksheet.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                if isinstance(cell.value, (int, float)):
                    cell.number_format = '#,##0.00'

        # Auto-adjust column widths
        for column in worksheet.columns:
            max_length = 0
            column_letter = get_column_letter(column[0].column)
            for cell in column:
                try:
                    if cell.value:
                        max_length = max(max_length, len(str(cell.value)))
                except:
                    pass
            adjusted_width = min(max_length + 2, 50)
            worksheet.column_dimensions[column_letter].width = adjusted_width
